"""Excel export of the whole ledger: Transactions, Accounts, People sheets."""

from io import BytesIO

from fastapi import APIRouter, Depends, Response
from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from backend.db import get_db
from backend.models import Transaction
from backend.services import accounts as accounts_svc
from backend.services import loans as loans_svc
from backend.utils import rupees, today_ist

router = APIRouter(prefix="/api", tags=["export"])

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _sheet(ws, headers: list[str], rows: list[list]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    for row in rows:
        ws.append(row)


@router.get("/export.xlsx")
def export_xlsx(db: Session = Depends(get_db)) -> Response:
    wb = Workbook()

    txs = db.scalars(
        select(Transaction)
        .options(
            selectinload(Transaction.account),
            selectinload(Transaction.to_account),
            selectinload(Transaction.category),
            selectinload(Transaction.person),
        )
        .order_by(Transaction.date, Transaction.id)
    ).all()
    ws = wb.active
    ws.title = "Transactions"
    _sheet(
        ws,
        ["id", "date", "type", "account", "to_account", "category", "person",
         "amount", "currency", "inr_amount", "note"],
        [
            [tx.id, tx.date, tx.type,
             tx.account.name if tx.account else None,
             tx.to_account.name if tx.to_account else None,
             tx.category.name if tx.category else None,
             tx.person.name if tx.person else None,
             rupees(tx.amount_minor), tx.currency, rupees(tx.inr_minor), tx.note]
            for tx in txs
        ],
    )
    for col in ("H", "J"):
        for cell in ws[col][1:]:
            cell.number_format = "#,##0.00"

    accounts = accounts_svc.list_with_balances(db, include_archived=True)
    _sheet(
        wb.create_sheet("Accounts"),
        ["name", "type", "currency", "balance", "outstanding", "archived"],
        [[a["name"], a["type"], a["currency"], a["balance"], a["outstanding"],
          a["is_archived"]] for a in accounts],
    )

    people = loans_svc.person_balances(db)
    ws_p = wb.create_sheet("People")
    _sheet(
        ws_p,
        ["name", "net (positive = owes me)", "last_activity"],
        [[p["name"], p["net"], p["last_activity"]] for p in people["persons"]],
    )
    ws_p.append([])
    ws_p.append(["Total owed to me", people["total_owed_to_me"]])
    ws_p.append(["Total I owe", people["total_i_owe"]])

    buf = BytesIO()
    wb.save(buf)
    return Response(
        content=buf.getvalue(),
        media_type=XLSX_MIME,
        headers={
            "Content-Disposition": f'attachment; filename="money-export-{today_ist()}.xlsx"'
        },
    )
